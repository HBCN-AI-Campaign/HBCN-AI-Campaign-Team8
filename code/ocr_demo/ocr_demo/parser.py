# -*- coding: utf-8 -*-
"""文档解析：Word / Excel / PDF 转文本。

- Word 用 python-docx（.docx）；旧 .doc 提示转换或忽略。
- Excel 用 openpyxl（.xlsx）。
- PDF 优先 pdfplumber 抽取文本层；抽取结果过少（扫描件）则转图走 OCR。
"""

from __future__ import annotations

from pathlib import Path

from . import config


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def parse_word(path: Path) -> str:
    if path.suffix.lower() == ".doc":
        return (
            "[提示] 旧版 .doc 格式暂不支持直接解析，"
            "请用 Word 另存为 .docx 后重试。"
        )
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []

    # 按文档顺序遍历段落与表格（保持原始阅读顺序）
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            text = para.text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            parts.append(_table_to_text(table))

    return "\n".join(parts)


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(" | ".join(cells))
    if not rows:
        return ""
    # 简单分隔：首行作表头，加一条分隔线
    header = rows[0]
    sep = " | ".join("---" for _ in header.split(" | "))
    return "\n".join([header, sep, *rows[1:]])


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def parse_excel(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return (
            "[提示] 旧版 .xls 格式暂不支持直接解析，"
            "请用 Excel 另存为 .xlsx 后重试。"
        )
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## 工作表：{ws.title}")
        for row in ws.iter_rows(values_only=True):
            # 跳过整行为空的行
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            cells = ["" if v is None else str(v).strip() for v in row]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def parse_pdf(path: Path, ocr_min_text: int = 20) -> tuple[str, bool]:
    """返回 (文本, 是否走了OCR)。文本层过少时视为扫描件走 OCR。"""
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                text_parts.append(t)

    text = "\n".join(text_parts).strip()
    # 文本层足够 → 直接返回
    if len(text) >= ocr_min_text:
        return text, False

    # 文本层过少（扫描件）→ 转图 OCR
    from .ocr_engine import recognize_image
    from .pdf_to_image import pdf_to_images

    ocr_lines: list[str] = []
    for img_path in pdf_to_images(path):
        ocr_lines.append(recognize_image(img_path))
    return "\n\n".join(ocr_lines), True


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------
def parse_file(path: Path) -> dict:
    """按文件类型分派解析，返回统一结构。"""
    kind = config.classify(path)
    result = {
        "path": str(path),
        "type": kind,
        "label": config.DOC_TYPE_LABELS.get(kind, "未知"),
        "text": "",
        "via_ocr": False,
        "error": None,
    }

    try:
        if kind == "image":
            from .ocr_engine import recognize_image

            result["text"] = recognize_image(path)
            result["via_ocr"] = True
        elif kind == "word":
            result["text"] = parse_word(path)
        elif kind == "excel":
            result["text"] = parse_excel(path)
        elif kind == "pdf":
            text, via_ocr = parse_pdf(path)
            result["text"] = text
            result["via_ocr"] = via_ocr
        else:
            result["error"] = f"不支持的文件类型：{path.suffix}"
    except Exception as exc:  # noqa: BLE001
        result["error"] = f"{type(exc).__name__}: {exc}"

    return result
