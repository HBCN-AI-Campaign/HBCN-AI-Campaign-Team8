#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成示例测试文档。

覆盖 5 类输入，用于验证 OCR Demo 的全链路：
  1. demo.docx            传统 Word（有文本层）
  2. demo.xlsx            Excel 表格（多 Sheet）
  3. demo_text.pdf        带文本层的 PDF（直接抽取）
  4. demo_scan.pdf        扫描件 PDF（整页只有图片，必须走 OCR）
  5. demo_scan.png        图片（走 OCR）

运行方式（uv 临时环境，不污染 Demo 主虚拟环境）：
    uv run --with python-docx --with openpyxl --with reportlab --with pillow scripts/make_samples.py
"""

from __future__ import annotations

import random
import sys
from pathlib import Path

SAMPLES = Path(__file__).resolve().parent.parent / "samples"
FONT_CJK = "C:/Windows/Fonts/msyh.ttc"
FONT_FALLBACK = "C:/Windows/Fonts/simhei.ttf"


# --------------------------------------------------------------------------
# 1. Word
# --------------------------------------------------------------------------
def make_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    title = doc.add_heading("境外贷款 · 授信申请函", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("致：XX 银行股份有限公司 国际业务部")
    doc.add_paragraph(
        "申请人：智汇跨境贸易有限公司（以下简称“申请人”），"
        "统一社会信用代码 91310000MA1KXXXXX，注册地为中国（上海）自由贸易试验区。"
    )
    doc.add_paragraph(
        "申请事项：申请开立金额为 USD 3,000,000.00（大写：叁佰万美元整）"
        "的境外流动资金贷款，期限 12 个月，用途为跨境采购原材料及支付物流费用。"
    )

    doc.add_heading("一、申请人基本信息", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "项目", "内容"
    for k, v in [
        ("公司名称", "智汇跨境贸易有限公司"),
        ("成立时间", "2016-04-18"),
        ("注册资本", "CNY 50,000,000.00"),
        ("法定代表人", "张伟"),
        ("主营业务", "跨境贸易、供应链管理"),
        ("上年度营业收入", "CNY 286,430,000.00"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = k, v

    doc.add_heading("二、风险提示", level=2)
    doc.add_paragraph(
        "申请人近 12 个月存在 2 笔跨境关联交易，交易对手注册于开曼群岛，"
        "需重点核查交易背景真实性及资金最终用途，防范洗钱及制裁合规风险。",
    )
    doc.add_paragraph(
        "本次贷款担保方式为：实际控制人夫妇连带责任保证 + 应收账款质押。",
    )

    p = doc.add_paragraph()
    run = p.add_run("备注：本文件为 OCR Demo 自动生成的测试样例，不代表真实业务。")
    run.font.size = Pt(9)
    run.italic = True

    doc.save(str(path))


# --------------------------------------------------------------------------
# 2. Excel
# --------------------------------------------------------------------------
def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")

    # Sheet1：财务数据
    ws = wb.active
    ws.title = "财务数据"
    rows = [
        ["科目", "2022年度", "2023年度", "2024年度", "单位"],
        ["营业收入", 218500000, 254300000, 286430000, "CNY"],
        ["营业成本", 176230000, 201450000, 224780000, "CNY"],
        ["毛利率", 19.34, 20.78, 21.52, "%"],
        ["净利润", 18340000, 22560000, 27090000, "CNY"],
        ["资产负债率", 48.2, 45.6, 43.1, "%"],
        ["经营性现金流", 12450000, 15870000, 19230000, "CNY"],
    ]
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 18
    for col in ("B", "C", "D", "E"):
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "A2"

    # Sheet2：跨境交易明细
    ws2 = wb.create_sheet("跨境交易明细")
    ws2.append(["交易日期", "交易对手", "币种", "金额", "国别", "用途"])
    for row in [
        ["2024-01-15", "SUNRISE TRADING PTE. LTD.", "USD", 480000, "新加坡", "采购原料"],
        ["2024-03-22", "GRAND OCEAN LOGISTICS LTD", "USD", 126500, "中国香港", "支付运费"],
        ["2024-05-09", "BLUE HORIZON HOLDINGS", "USD", 750000, "开曼群岛", "设备采购"],
        ["2024-07-30", "NORDIC PAPER AB", "EUR", 218000, "瑞典", "采购原料"],
        ["2024-09-11", "PACIFIC FREIGHT CO., LTD.", "USD", 98400, "日本", "支付运费"],
    ]:
        ws2.append(row)
    for cell in ws2[1]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDEF", (13, 30, 8, 14, 12, 12)):
        ws2.column_dimensions[col].width = w

    wb.save(str(path))


# --------------------------------------------------------------------------
# 3. 带文本层的 PDF
# --------------------------------------------------------------------------
def make_text_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font_path = FONT_CJK if Path(FONT_CJK).exists() else FONT_FALLBACK
    pdfmetrics.registerFont(TTFont("CJK", font_path))

    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setTitle("境外贷款支持性文件清单")
    c.setFont("CJK", 18)
    c.drawString(60, h - 90, "境外贷款 · 支持性文件清单")

    c.setFont("CJK", 10.5)
    lines = [
        "文件编号：HBCN-2024-CR-0871",
        "客户名称：智汇跨境贸易有限公司",
        "业务类型：境外流动资金贷款",
        "拟授信金额：USD 3,000,000.00",
        "",
        "一、基础证照类",
        "    1. 营业执照（统一社会信用代码 91310000MA1KXXXXX）",
        "    2. 公司章程及最新工商变更登记",
        "    3. 法定代表人身份证及授权委托书",
        "",
        "二、财务资料类",
        "    4. 近三年审计报告（2022 / 2023 / 2024）",
        "    5. 最近一期财务报表及银行流水",
        "    6. 增值税纳税申报表",
        "",
        "三、交易背景类",
        "    7. 境外采购合同（SUNRISE TRADING PTE. LTD.）",
        "    8. 物流运输单据及报关单",
        "    9. 跨境付款路径说明",
        "",
        "四、合规提示",
        "    交易对手涉及开曼群岛注册主体，须执行加强型尽职调查（EDD），",
        "    核查最终受益所有人（UBO）并比对制裁名单与政治敏感人物（PEP）名单。",
    ]
    y = h - 140
    for line in lines:
        c.drawString(60, y, line)
        y -= 22
    c.setFont("CJK", 9)
    c.drawString(60, 60, "注：本文件为 OCR Demo 自动生成的测试样例，不代表真实业务。")
    c.showPage()
    c.save()


# --------------------------------------------------------------------------
# 4/5. 扫描件（先出 PNG，再由 PNG 生成图片型 PDF）
# --------------------------------------------------------------------------
def render_scan_image() -> "Image.Image":  # noqa: F821
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    random.seed(20241019)
    W, H = 1240, 1754  # A4 @150dpi
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)

    font_path = FONT_CJK if Path(FONT_CJK).exists() else FONT_FALLBACK
    f_title = ImageFont.truetype(font_path, 34)
    f_body = ImageFont.truetype(font_path, 22)
    f_small = ImageFont.truetype(font_path, 17)

    draw.rectangle([40, 40, W - 40, H - 40], outline="#333333", width=2)
    draw.text((80, 90), "境 外 贷 款 支 持 性 文 件", font=f_title, fill="#111111")
    draw.line([80, 145, W - 80, 145], fill="#333333", width=2)

    body = [
        ("文件编号：", "HBCN-2024-CR-0871"),
        ("客户名称：", "智汇跨境贸易有限公司"),
        ("业务类型：", "境外流动资金贷款（跨境采购）"),
        ("拟授信金额：", "USD 3,000,000.00"),
        ("贷款期限：", "12 个月"),
        ("担保方式：", "实际控制人保证 + 应收账款质押"),
        ("", ""),
        ("风险提示：", "交易对手含开曼群岛注册主体，需执行加强型尽调"),
        ("合规提示：", "需核查 UBO 并比对制裁名单与 PEP 名单"),
        ("", ""),
        ("经办人：", "李娜    复核人：王强    日期：2024-10-19"),
    ]
    y = 190
    for label, value in body:
        if label:
            draw.text((95, y), label, font=f_body, fill="#444444")
            draw.text((95 + 165, y), value, font=f_body, fill="#101010")
        y += 46

    # 手绘一个表格，测试 OCR 对表格结构的识别
    ty = y + 30
    draw.text((95, ty - 40), "近三年主要财务数据", font=f_body, fill="#101010")
    cols = [95, 300, 560, 820, 1080]
    headers = ["科目", "2022年度", "2023年度", "2024年度"]
    for i, htxt in enumerate(headers):
        draw.text((cols[i] + 12, ty + 10), htxt, font=f_small, fill="#111111")
    draw.rectangle([cols[0], ty, cols[-1], ty + 42], outline="#555555", width=2)
    rows = [
        ["营业收入", "218,500,000", "254,300,000", "286,430,000"],
        ["净利润", "18,340,000", "22,560,000", "27,090,000"],
        ["资产负债率", "48.2%", "45.6%", "43.1%"],
    ]
    for r, row in enumerate(rows):
        ry = ty + 42 + r * 42
        for ci, txt in enumerate(row):
            draw.text((cols[ci] + 12, ry + 10), txt, font=f_small, fill="#111111")
        draw.rectangle([cols[0], ry, cols[-1], ry + 42], outline="#888888", width=1)
    for x in cols:
        draw.line([x, ty, x, ty + 42 * (len(rows) + 1)], fill="#888888", width=1)

    draw.text(
        (95, H - 120),
        "本文件为 OCR Demo 自动生成的扫描件模拟样例，不代表真实业务。",
        font=f_small,
        fill="#666666",
    )

    # 模拟扫描件：轻微倾斜 + 模糊 + 噪点
    img = img.rotate(-0.35, resample=Image.BICUBIC, fillcolor="white")
    img = img.filter(ImageFilter.GaussianBlur(radius=0.6))
    px = img.load()
    for _ in range(40000):
        x, y = random.randrange(W), random.randrange(H)
        v = 255 - random.randint(0, 28)
        px[x, y] = (v, v, v)
    return img


def main() -> int:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    made = []

    make_docx(SAMPLES / "demo.docx")
    made.append("demo.docx")

    make_xlsx(SAMPLES / "demo.xlsx")
    made.append("demo.xlsx")

    make_text_pdf(SAMPLES / "demo_text.pdf")
    made.append("demo_text.pdf")

    img = render_scan_image()
    img.save(SAMPLES / "demo_scan.png")
    made.append("demo_scan.png")

    img.save(SAMPLES / "demo_scan.pdf", "PDF", resolution=150.0)
    made.append("demo_scan.pdf")

    print(f"[OK] 示例文件已生成于 {SAMPLES}")
    for name in made:
        p = SAMPLES / name
        print(f"     - {name:<18} {p.stat().st_size / 1024:8.1f} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
