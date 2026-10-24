# -*- coding: utf-8 -*-
"""文档解析：Word / Excel / PDF 转文本。

- Word 用 python-docx（.docx）；旧 .doc 提示转换或忽略。
- Excel 用 openpyxl（.xlsx）。
- PDF 优先 pdfplumber 抽取文本层；抽取结果过少（扫描件）则转图走 OCR。
"""

from __future__ import annotations

import logging
import shutil
from pathlib import Path

from . import config
from .vision_ocr import is_configured, recognize_page

log = logging.getLogger(__name__)


def _recognize_image_detailed(image_path: Path, page: int, min_score: float) -> tuple[list[dict], str]:
    from .ocr_engine import recognize_image_detailed

    local_items = recognize_image_detailed(image_path, min_score=min_score)
    scores = [float(item.get("score", 0.0)) for item in local_items]
    should_fallback = is_configured() and (
        len(local_items) < config.OCR_VISION_TRIGGER_MIN_LINES
        or (scores and sum(scores) / len(scores) < config.OCR_VISION_TRIGGER_MIN_SCORE)
    )
    if should_fallback:
        text = recognize_page(image_path)
        if text:
            candidate_items = [
                {"text": line, "score": 0.0, "box": None, "poly": None,
                 "page": page, "method": "vision_ocr"}
                for line in text.splitlines() if line.strip()
            ]
            if candidate_items:
                return candidate_items, "vision_ocr"
    return [dict(item, page=page, method="paddleocr") for item in local_items], "paddleocr"


def _remove_rendered_pages(images: list[Path]) -> None:
    if images:
        shutil.rmtree(images[0].parent, ignore_errors=True)


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def parse_word(path: Path) -> str:
    if path.suffix.lower() == ".doc":
        return (
            "[提示] 旧版 .doc 格式暂不支持直接解析，"
            "请用 Word 另存为 .docx 后重试。"
        )
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []

    # 按文档顺序遍历段落与表格（保持原始阅读顺序）
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            text = para.text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            parts.append(_table_to_text(table))

    return "\n".join(parts)


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(" | ".join(cells))
    if not rows:
        return ""
    # 简单分隔：首行作表头，加一条分隔线
    header = rows[0]
    sep = " | ".join("---" for _ in header.split(" | "))
    return "\n".join([header, sep, *rows[1:]])


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def parse_excel(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return (
            "[提示] 旧版 .xls 格式暂不支持直接解析，"
            "请用 Excel 另存为 .xlsx 后重试。"
        )
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## 工作表：{ws.title}")
        for row in ws.iter_rows(values_only=True):
            # 跳过整行为空的行
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            cells = ["" if v is None else str(v).strip() for v in row]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _parse_pdf_with_method(path: Path, ocr_min_text: int = 20) -> tuple[str, bool, str]:
    """Return text, OCR usage, and the actual extraction method."""
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                text_parts.append(t)

    text = "\n".join(text_parts).strip()
    # 文本层足够 → 直接返回
    if len(text) >= ocr_min_text:
        return text, False, "native_pdf"

    # 文本层过少（扫描件）→ 转图 OCR
    from .pdf_to_image import pdf_to_images

    images = pdf_to_images(path)
    ocr_lines: list[str] = []
    methods: set[str] = set()
    try:
        for page_number, img_path in enumerate(images, start=1):
            items, method = _recognize_image_detailed(img_path, page_number, 0.3)
            methods.add(method)
            ocr_lines.append("\n".join(item["text"] for item in items))
    finally:
        _remove_rendered_pages(images)
    method = "vision_ocr" if "vision_ocr" in methods else "paddleocr"
    return "\n\n".join(ocr_lines), True, method


def parse_pdf(path: Path, ocr_min_text: int = 20) -> tuple[str, bool]:
    """Return (text, whether OCR was used), preserving the original API."""
    text, via_ocr, _ = _parse_pdf_with_method(path, ocr_min_text)
    return text, via_ocr


def parse_image_detailed(path: Path, min_score: float = 0.3) -> dict:
    """识别图片，返回带坐标的明细结果。"""
    items, method = _recognize_image_detailed(path, 1, min_score)
    text = "\n".join(it["text"] for it in items)
    return {"text": text, "blocks": items, "via_ocr": True, "method": method}


def parse_pdf_detailed(
    path: Path, ocr_min_text: int = 20, min_score: float = 0.3
) -> dict:
    """解析 PDF，返回带坐标的明细结果。

    文本型 PDF：用 pdfplumber 抽取词级坐标（每页的 words）。
    扫描件 PDF：逐页转图走 OCR，返回 rec_boxes/rec_polys 坐标。
    """
    import pdfplumber

    pages_text: list[str] = []
    pages: list = []
    with pdfplumber.open(str(path)) as pdf:
        for pi, page in enumerate(pdf.pages, start=1):
            t = page.extract_text() or ""
            pages_text.append(t)
            pages.append(page)

    total_text = "\n".join(pages_text).strip()
    if len(total_text) >= ocr_min_text:
        # 文本型：取词级坐标（每个词：text/x0/top/x1/bottom）
        blocks: list[dict] = []
        for pi, page in enumerate(pages, start=1):
            for w in page.extract_words():
                blocks.append(
                    {
                        "text": w["text"],
                        "score": 1.0,
                        "box": [w["x0"], w["top"], w["x1"], w["bottom"]],
                        "poly": [
                            [w["x0"], w["top"]],
                            [w["x1"], w["top"]],
                            [w["x1"], w["bottom"]],
                            [w["x0"], w["bottom"]],
                        ],
                        "page": pi,
                        "method": "native_pdf",
                    }
                )
        return {"text": total_text, "blocks": blocks, "via_ocr": False, "method": "native_pdf"}

    # 扫描件：转图 OCR，按页给坐标
    from .pdf_to_image import pdf_to_images

    all_blocks: list[dict] = []
    all_text: list[str] = []
    imgs = pdf_to_images(path)
    try:
        for pi, img_path in enumerate(imgs, start=1):
            items, _ = _recognize_image_detailed(img_path, pi, min_score)
            all_blocks.extend(items)
            all_text.append("\n".join(it["text"] for it in items))
    finally:
        _remove_rendered_pages(imgs)
    return {
        "text": "\n\n".join(all_text),
        "blocks": all_blocks,
        "via_ocr": True,
        "method": "vision_ocr" if any(item.get("method") == "vision_ocr" for item in all_blocks) else "paddleocr",
    }


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------
def parse_file(path: Path, detailed: bool = False) -> dict:
    """按文件类型分派解析，返回统一结构。

    detailed=True 时，图片/PDF 额外返回 `blocks`（带坐标的明细）。
    """
    kind = config.classify(path)
    result = {
        "path": str(path),
        "type": kind,
        "label": config.DOC_TYPE_LABELS.get(kind, "未知"),
        "text": "",
        "via_ocr": False,
        "method": "native",
        "error": None,
    }

    try:
        if kind == "image":
            if detailed:
                res = parse_image_detailed(path)
                result["text"] = res["text"]
                result["via_ocr"] = True
                result["blocks"] = res["blocks"]
                result["method"] = res["method"]
            else:
                res = parse_image_detailed(path)
                result["text"] = res["text"]
                result["via_ocr"] = True
                result["method"] = res["method"]
        elif kind == "word":
            result["text"] = parse_word(path)
        elif kind == "excel":
            result["text"] = parse_excel(path)
        elif kind == "pdf":
            if detailed:
                res = parse_pdf_detailed(path)
                result["text"] = res["text"]
                result["via_ocr"] = res["via_ocr"]
                result["blocks"] = res["blocks"]
                result["method"] = res["method"]
            else:
                text, via_ocr, method = _parse_pdf_with_method(path)
                result["text"] = text
                result["via_ocr"] = via_ocr
                result["method"] = method
        else:
            result["error"] = f"不支持的文件类型：{path.suffix}"
    except Exception as exc:  # noqa: BLE001
        result["error_type"] = type(exc).__name__
        log.error("Document parsing failed (%s)", result["error_type"])
        result["error"] = f"{type(exc).__name__}: {exc}"

    return result
