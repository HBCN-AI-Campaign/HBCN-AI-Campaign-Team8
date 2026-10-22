from __future__ import annotations

import io
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from .config import settings
from .ocr import OCRClient, _read_text


class ParseError(RuntimeError):
    pass


class HtmlText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip += 1
        if not self.skip and tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def extract_pages(filename: str, content: bytes) -> list[dict]:
    extension = Path(filename).suffix.lower()
    pages = []
    if extension == ".pdf":
        from pypdf import PdfReader

        document = PdfReader(io.BytesIO(content))
        if document.is_encrypted:
            raise ParseError("文件有密码保护，请上传解密版本")
        if len(document.pages) > settings.RAG_MAX_DOCUMENT_PAGES:
            raise ParseError("文档页数超过处理上限，请拆分后上传")
        total = 0
        for index, page in enumerate(document.pages):
            text = (page.extract_text() or "").strip()
            method = "native_pdf"
            if not text:
                import pypdfium2

                with pypdfium2.PdfDocument(content) as rendered:
                    rendered_page = rendered[index]
                    if rendered_page.get_width() * rendered_page.get_height() > 20_000_000:
                        raise ParseError("扫描页尺寸过大，请压缩页面后上传")
                    bitmap = rendered_page.render(scale=1.5)
                    image = bitmap.to_pil()
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    image.close()
                    bitmap.close()
                    rendered_page.close()
                result = OCRClient().extract(f"page-{index + 1}.png", buffer.getvalue())
                if result["error"]:
                    raise ParseError("扫描页识别失败，请确认 OCR 或视觉模型可用")
                text = result["text"].strip()
                method = result.get("method", "ocr_unknown")
            total += len(text)
            if total > settings.RAG_MAX_TEXT_CHARS:
                raise ParseError("文档文本超过处理上限，请拆分后上传")
            reasons = ["pdf_images_or_ocr_layer_unverified"] if method == "native_pdf" and len(page.images) else []
            pages.append({"page": index + 1, "text": text, "extraction": {"method": method, "review_reasons": reasons}})
    elif extension in {".docx", ".xlsx", ".pptx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 150 * 1024 * 1024:
                    raise ParseError("文档解压后过大，请拆分后上传")
        except zipfile.BadZipFile as exc:
            raise ParseError("文件格式无效") from exc
        if extension == ".docx":
            from docx import Document
            from docx.table import Table

            document = Document(io.BytesIO(content))
            parts = []
            for item in document.iter_inner_content():
                if isinstance(item, Table):
                    parts.extend(" | ".join(cell.text for cell in row.cells) for row in item.rows)
                else:
                    parts.append(item.text)
            reasons = []
            if document.inline_shapes:
                reasons.append("embedded_images_not_transcribed")
            if any(paragraph.text.strip() for section in document.sections for container in (section.header, section.footer) for paragraph in container.paragraphs):
                reasons.append("header_footer_not_transcribed")
            pages = [{"page": None, "text": "\n".join(parts), "extraction": {"method": "docx_native", "review_reasons": reasons}}]
        elif extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            formulas = None
            try:
                formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                for sheet in workbook:
                    lines, count, missing = [], 0, False
                    for values, expression_row in zip(sheet.iter_rows(values_only=True), formulas[sheet.title].iter_rows(values_only=True)):
                        missing = missing or any(isinstance(expression, str) and expression.startswith("=") and value is None
                            for value, expression in zip(values, expression_row))
                        line = " | ".join(str(value) if value is not None else "" for value in values)
                        count += len(line) + 1
                        if count > settings.RAG_MAX_TEXT_CHARS:
                            raise ParseError("工作表文本超过处理上限，请拆分后上传")
                        lines.append(line)
                    pages.append({"page": None, "section": sheet.title, "text": "\n".join(lines),
                        "extraction": {"method": "xlsx_cell_values", "review_reasons": ["formula_values_missing"] if missing else []}})
            finally:
                if formulas is not None:
                    formulas.close()
                workbook.close()
        else:
            raise ParseError("当前解析器暂不支持 PowerPoint，请转换为 PDF")
    elif extension in {".txt", ".md", ".csv", ".json", ".mmd", ".html", ".htm"}:
        text = _read_text(content)
        reasons = []
        if not content.startswith((b"\xff\xfe", b"\xfe\xff", b"\x00\x00\xfe\xff")):
            try:
                content.decode("utf-8-sig")
            except UnicodeDecodeError:
                reasons.append("encoding_inferred")
        if extension in {".html", ".htm"}:
            parser = HtmlText()
            parser.feed(text)
            text = "".join(parser.parts)
        pages = [{"page": None, "text": text, "extraction": {"method": "text_decode", "review_reasons": reasons}}]
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}:
        result = OCRClient().extract(filename, content)
        if result["error"]:
            raise ParseError("图片识别失败，请确认 OCR 服务可用")
        pages = [{"page": 1, "text": result["text"], "extraction": {"method": result.get("method", "ocr_unknown")}}]
    else:
        raise ParseError("当前解析器不支持此格式，请转换为 PDF、DOCX、XLSX 或文本")
    total = sum(len(page["text"]) for page in pages)
    if total > settings.RAG_MAX_TEXT_CHARS:
        raise ParseError("文档文本超过处理上限，请拆分后上传")
    if not any(page["text"].strip() for page in pages):
        raise ParseError("没有提取到可检索文字")
    return pages
