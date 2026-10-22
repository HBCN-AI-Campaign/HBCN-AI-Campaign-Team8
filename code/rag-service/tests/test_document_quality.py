import io
import unittest
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import Workbook

from app.document_quality import assess_pages
from app.ocr import OCRClient, _read_text
from app.parsing import extract_pages


class DocumentQualityTests(unittest.TestCase):
    def test_mechanical_checks_preserve_multilingual_facts_without_claiming_accuracy(self):
        text = "支付指令 USD 5,000.00 / 金额不符\nPayment amount discrepancy; العربية; € 1.23"
        report = assess_pages([{"page": None, "text": text, "extraction": {"method": "native_pdf"}}])
        self.assertEqual(report["status"], "mechanically_checked")
        self.assertIn("not_ocr_accuracy", report["limits"])
        self.assertNotIn(text, str(report))
        for invalid in ["金额\ufffd错误", "USD 50\x000", "bad\ud800text"]:
            self.assertEqual(assess_pages([{"page": 1, "text": invalid}])["status"], "quarantined")
        changed = assess_pages([{"page": None, "text": text.replace("5,000", "500")}])
        self.assertNotEqual(changed["parsed_sha256"], report["parsed_sha256"])

    def test_ocr_empty_and_unmapped_units_require_review_without_fabricated_confidence(self):
        pages = [{"page": 1, "text": "完整文字", "extraction": {"method": "ocr_service"}},
                 {"page": 2, "text": ""}, {"page": 3, "text": "(cid:12)"}]
        report = assess_pages(pages)
        self.assertEqual(report["status"], "quality_review")
        self.assertEqual(report["unit_count"], 3)
        self.assertEqual(set(report["review_reasons"]), {"ocr_requires_source_comparison", "empty_unit", "unmapped_glyphs", "extraction_provenance_missing"})
        self.assertNotIn("confidence", report)

    def test_unicode_boms_decode_exactly_and_corrupt_bytes_remain_detectable(self):
        text = "金额 5000 USD"
        for codec in ["utf-8-sig", "utf-16", "utf-32", "gb18030"]:
            self.assertEqual(_read_text(text.encode(codec)), text)
        self.assertEqual(assess_pages([{"page": None, "text": _read_text(b'\xff')}])["status"], "quarantined")

    def test_visual_ocr_requires_stop_and_plain_text_and_hides_provider_errors(self):
        client = OCRClient()
        for finish in ["length", "content_filter", None]:
            response = Mock(status_code=200)
            response.json.return_value = {"choices": [{"finish_reason": finish, "message": {"content": "金额500"}}]}
            with patch("app.ocr.httpx.post", return_value=response):
                result = client._call_vision("test.png", b"synthetic", ".png")
            self.assertTrue(result["error"])
            self.assertEqual(result["text"], "")
        response.json.return_value = {"choices": [{"finish_reason": "stop", "message": {"content": "金额5000"}}]}
        with patch("app.ocr.httpx.post", return_value=response):
            self.assertEqual(client._call_vision("test.png", b"synthetic", ".png")["text"], "金额5000")
        response.status_code, response.text = 503, "SYNTHETIC-SECRET"
        with patch("app.ocr.httpx.post", return_value=response):
            self.assertNotIn("SYNTHETIC-SECRET", str(client._call_vision("test.png", b"synthetic", ".png")))

    def test_docx_retains_table_position_and_flags_untranscribed_headers(self):
        document = Document()
        document.add_paragraph("表前条款")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "金额 5000 USD"
        document.add_paragraph("表后例外")
        document.sections[0].header.paragraphs[0].text = "重要适用条件"
        content = io.BytesIO()
        document.save(content)
        pages = extract_pages("test.docx", content.getvalue())
        self.assertLess(pages[0]["text"].index("5000"), pages[0]["text"].index("表后"))
        self.assertIn("header_footer_not_transcribed", assess_pages(pages)["review_reasons"])

    def test_xlsx_missing_formula_cache_is_not_an_empty_verified_amount(self):
        book = Workbook()
        book.active.append(["金额", "USD", "=5000+5"])
        content = io.BytesIO()
        book.save(content)
        book.close()
        pages = extract_pages("test.xlsx", content.getvalue())
        self.assertIn("formula_values_missing", assess_pages(pages)["review_reasons"])
        self.assertNotIn("5005", pages[0]["text"])
