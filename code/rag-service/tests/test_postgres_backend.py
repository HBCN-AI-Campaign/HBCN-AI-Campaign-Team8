import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import httpx
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from openpyxl import Workbook
from fastapi.testclient import TestClient

from app.answer_graph import build_answer_graph
from app.backend_api import AnswerRequest
from app.config import settings
from app.main import app
from app.parsing import ParseError, extract_pages
from app.postgres_engine import BailianEmbeddings
from app.runtime_client import RuntimeErrorResponse, SDKRuntime


def sample_evidence():
    return [{"content": "贷款资金不得用于股权投资", "document_id": "doc-1", "document_keyword": "业务规则.xlsx",
             "page_num": None, "similarity": 0.8, "chunk_id": "chunk-1", "metadata": {}}]


class ParsingTests(unittest.TestCase):
    def test_pdf_preserves_real_pages(self):
        buffer = io.BytesIO()
        document = Canvas(buffer)
        document.drawString(50, 700, "First source page")
        document.showPage()
        document.drawString(50, 700, "Second source page")
        document.save()
        pages = extract_pages("sample.pdf", buffer.getvalue())
        self.assertEqual([page["page"] for page in pages], [1, 2])
        self.assertIn("Second source", pages[1]["text"])

    def test_html_does_not_index_scripts(self):
        pages = extract_pages("source.html", b"<h1>Rule</h1><script>secret()</script><p>Evidence</p>")
        self.assertIn("Evidence", pages[0]["text"])
        self.assertNotIn("secret", pages[0]["text"])
        self.assertIsNone(pages[0]["page"])

    def test_docx_has_no_fabricated_page(self):
        document = Document()
        document.add_paragraph("真实业务规则")
        buffer = io.BytesIO()
        document.save(buffer)
        pages = extract_pages("rules.docx", buffer.getvalue())
        self.assertIn("真实业务规则", pages[0]["text"])
        self.assertIsNone(pages[0]["page"])

    def test_empty_and_oversized_text_are_rejected(self):
        with self.assertRaises(ParseError):
            extract_pages("empty.txt", b" ")
        with patch.object(settings, "RAG_MAX_TEXT_CHARS", 1000), self.assertRaises(ParseError):
            extract_pages("large.txt", b"x" * 1001)

    def test_office_zip_bomb_guard(self):
        with self.assertRaises(ParseError):
            extract_pages("invalid.docx", b"invalid zip")

    def test_bad_embedding_dimensions_are_rejected(self):
        with patch("app.postgres_engine.EmbeddingClient") as client:
            client.return_value.embed.return_value = [[0.1]]
            with self.assertRaisesRegex(Exception, "维度"):
                BailianEmbeddings().embed_documents(["test"])

    def test_scanned_pdf_uses_renderer_and_ocr_without_inventing_text(self):
        buffer = io.BytesIO()
        document = Canvas(buffer)
        document.showPage()
        document.save()
        with patch("app.parsing.OCRClient") as client:
            client.return_value.extract.return_value = {"text": "扫描原文", "error": None}
            pages = extract_pages("scan.pdf", buffer.getvalue())
            self.assertEqual(pages[0]["page"], 1)
            self.assertEqual(pages[0]["text"], "扫描原文")
            self.assertEqual(pages[0]["extraction"]["method"], "ocr_unknown")
            self.assertTrue(client.return_value.extract.call_args.args[1].startswith(b"\x89PNG"))

    def test_xlsx_preserves_worksheet_not_fabricated_page(self):
        workbook = Workbook()
        workbook.active.title = "贷款用途"
        workbook.active.append(["项目", "用途"])
        workbook.active.append(["贷款", "营运资金"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        pages = extract_pages("rules.xlsx", buffer.getvalue())
        self.assertEqual(pages[0]["section"], "贷款用途")
        self.assertIn("营运资金", pages[0]["text"])
        self.assertIsNone(pages[0]["page"])


class AnswerGraphTests(unittest.TestCase):
    def state(self):
        return {"query": "贷款用途有什么限制", "scope": "all", "top_k": 5, "attempts": 0, "warnings": []}

    def test_no_evidence_plans_and_checks_but_does_not_generate(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = []
        runtime.complete.side_effect = ['{"intent":"rule","queries":[]}', '{"relevant":false,"sufficient":false}']
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertEqual(runtime.complete.call_count, 2)
        self.assertEqual(result["grounding"], "insufficient")
        self.assertFalse(result["sources"])

    def test_answer_uses_real_sources(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = sample_evidence()
        runtime.complete.side_effect = ['{"intent":"rule","queries":[]}', '{"relevant":true,"sufficient":true,"selected":[1],"rewrite":""}', '{"statements": [{"text": "不得用于股权投资。", "supports": [{"citation": 1, "quote": "贷款资金不得用于股权投资"}]}]}',
            '{"checks":[{"segment":1,"verdict":"supported","supports":[{"citation":1,"quote":"贷款资金不得用于股权投资"}],"reason":"禁止用途一致"}]}']
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertEqual(result["grounding"], "cited")
        self.assertEqual(result["sources"][0]["document_id"], "doc-1")

    def test_invalid_citations_are_not_presented_as_supported_answer(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = sample_evidence()
        runtime.complete.side_effect = ['{"intent":"rule","queries":[]}', '{"relevant":true,"sufficient":true,"selected":[1]}', '{"statements": [{"text": "伪造的结论", "supports": [{"citation": 9, "quote": "贷款资金不得用于股权投资"}]}]}']
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertNotIn("伪造的结论", result["answer"])
        self.assertEqual(result["grounding"], "insufficient")

    def test_retrieval_refinement_is_bounded_at_two_rounds(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = sample_evidence()
        runtime.complete.side_effect = ['{"intent":"rule","queries":[]}',
            '{"relevant":false,"sufficient":false,"selected":[],"rewrite":"更准确的贷款用途限制"}',
            '{"relevant":false,"sufficient":false,"selected":[]}']
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertEqual(result["retrieval_rounds"], 2)
        self.assertEqual(engine.retrieve.call_count, 2)
        self.assertEqual(runtime.complete.call_count, 2)

    def test_malformed_assessment_abstains(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = sample_evidence()
        runtime.complete.return_value = "不是JSON"
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertEqual(result["grounding"], "insufficient")

    def test_conversation_rewrites_search_but_is_not_evidence(self):
        engine, runtime = Mock(), Mock()
        engine.retrieve.return_value = []
        runtime.complete.side_effect = [
            '{"intent":"rule","queries":[{"text":"固定资产贷款用途限制","family":"general"}]}',
            '{"relevant":false,"sufficient":false}',
        ]
        state = {**self.state(), "conversation": [{"role": "user", "content": "固定资产贷款"}]}
        result = build_answer_graph(engine, runtime).invoke(state)["result"]
        self.assertEqual(engine.retrieve.call_args.kwargs["query"], "固定资产贷款用途限制")
        self.assertFalse(result["sources"])

    def test_law_evidence_is_always_qualified(self):
        engine, runtime = Mock(), Mock()
        evidence = sample_evidence()
        evidence[0]["metadata"] = {"scope": "law"}
        engine.retrieve.return_value = evidence
        runtime.complete.side_effect = ['{"intent":"law","queries":[]}', '{"relevant":true,"sufficient":true,"selected":[1]}', '{"statements": [{"text": "贷款资金不得用于股权投资。", "supports": [{"citation": 1, "quote": "贷款资金不得用于股权投资"}]}]}',
            '{"checks":[{"segment":1,"verdict":"supported","supports":[{"citation":1,"quote":"贷款资金不得用于股权投资"}],"reason":"原文有此限制"}]}']
        result = build_answer_graph(engine, runtime).invoke(self.state())["result"]
        self.assertEqual(result["grounding"], "cited")
        self.assertTrue(any("效力" in warning for warning in result["warnings"]))

    def test_task_documents_are_not_accepted_by_ask_schema(self):
        with self.assertRaises(ValueError):
            AnswerRequest(query="规则", scope="task_documents")

    def test_stream_reports_graph_steps_and_final_answer(self):
        engine = Mock()
        engine.retrieve.return_value = []
        events = list(build_answer_graph(engine, Mock()).stream(self.state(), stream_mode=["custom", "updates"]))
        self.assertTrue(any(mode == "custom" and value["status"] == "running" for mode, value in events))
        self.assertTrue(any(mode == "updates" and "generate" in value for mode, value in events))


class RuntimeTests(unittest.TestCase):
    def call_runtime(self, body, status_code=200):
        transport = httpx.MockTransport(lambda request: httpx.Response(status_code, text=body))
        client = httpx.AsyncClient(transport=transport)
        with patch("app.runtime_client.httpx.AsyncClient", return_value=client), patch.object(settings, "ASSISTANT_RUNTIME_URL", "http://runtime"), patch("app.runtime_client.runtime_token", return_value="test-only"):
            return SDKRuntime().complete(message="问题", instructions="要求", evidence=[])

    def test_reads_sdk_final_result_not_steps(self):
        result = self.call_runtime('event: step\ndata: {"event":"step"}\n\nevent: assistant_result\ndata: {"event":"assistant_result","answer":"真实回答"}\n\n')
        self.assertEqual(result, "真实回答")

    def test_failed_sdk_response_is_not_faked(self):
        with self.assertRaises(RuntimeErrorResponse):
            self.call_runtime('data: {"event":"assistant_result","answer":null}\n\n')
        with self.assertRaises(RuntimeErrorResponse) as error:
            self.call_runtime("secret upstream detail", 500)
        self.assertNotIn("secret upstream detail", str(error.exception))


class BackendContractTests(unittest.TestCase):
    def setUp(self):
        self.token_patch = patch.object(settings, "SERVICE_TOKEN", "test-only-service-token")
        self.token_patch.start()
        self.addCleanup(self.token_patch.stop)
        self.client = TestClient(app)
        self.headers = {"Authorization": "Bearer test-only-service-token"}

    def test_new_endpoints_require_auth(self):
        for path in ["/api/v1/ready", "/api/v1/documents/not-a-uuid/content", "/api/v1/documents/not-a-uuid/source"]:
            self.assertEqual(self.client.get(path).status_code, 401)
        self.assertEqual(self.client.post("/api/v1/answers", json={"query": "test"}).status_code, 401)
        self.assertEqual(self.client.post("/api/v1/retrievals", json={"query": "test"}).status_code, 401)

    def test_task_scoped_query_and_empty_questions_are_rejected(self):
        for query in [{"query": " "}, {"query": "question", "scope": "task_documents"}]:
            self.assertEqual(self.client.post("/api/v1/answers", headers=self.headers, json=query).status_code, 422)

    def test_ingestion_does_not_claim_ready_before_processing(self):
        engine = Mock()
        engine.submit.return_value = ({"id": "doc1", "state": "queued"}, True)
        with patch("app.backend_api.postgres_engine", return_value=engine):
            response = self.client.post("/api/v1/ingestions", headers=self.headers, files={"file": ("rules.txt", b"rules")})
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.json()["document"]["state"], "queued")

    def test_readiness_reports_legacy_quality_backlog_as_degraded(self):
        engine = Mock()
        engine.readiness.return_value = {"unverified_ready_documents": 2, "quality_policy_version": "document-quality-v1"}
        with patch("app.backend_api.postgres_engine", return_value=engine), patch("app.backend_api.SDKRuntime") as runtime:
            runtime.return_value.health.return_value = True
            response = self.client.get("/api/v1/ready", headers=self.headers)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["status"], "degraded")
        self.assertEqual(response.json()["unverified_ready_documents"], 2)

    def test_empty_unsupported_and_oversized_files_are_rejected(self):
        for filename, content, status in [("old.doc", b"test", 415), ("empty.txt", b"", 422), ("big.txt", b"x" * 20, 413)]:
            with patch.object(settings, "MAX_UPLOAD_BYTES", 10):
                response = self.client.post("/api/v1/ingestions", headers=self.headers, files={"file": (filename, content)})
            self.assertEqual(response.status_code, status)

    def test_source_url_must_not_contain_credentials(self):
        response = self.client.post("/api/v1/ingestions", headers=self.headers,
                                    data={"source_url": "https://user:secret@example.com/doc"},
                                    files={"file": ("rules.txt", b"rules")})
        self.assertEqual(response.status_code, 422)

    def test_original_active_content_is_attachment_and_nosniff(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "original"
            path.write_bytes(b"<script>alert(1)</script>")
            engine = Mock()
            engine.documents.get.return_value = {"name": "unsafe.html", "state": "ready"}
            engine.documents.original_path.return_value = path
            with patch("app.backend_api.postgres_engine", return_value=engine):
                response = self.client.get("/api/v1/documents/11111111-1111-4111-8111-111111111111/source", headers=self.headers)
            self.assertEqual(response.content, path.read_bytes())
            self.assertIn("attachment", response.headers["content-disposition"])
            self.assertEqual(response.headers["x-content-type-options"], "nosniff")


if __name__ == "__main__":
    unittest.main()
